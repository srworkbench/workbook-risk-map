"""Read XLSX formula dependencies without evaluating formulas or changing the workbook."""
from __future__ import annotations
import argparse
from collections import defaultdict, deque
import hashlib
import io
import json
from pathlib import Path
import re
import shutil
import zipfile
from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import TokenizerError
from openpyxl.utils.cell import range_boundaries, get_column_letter

MAX_RANGE = 2000
MAX_FORMULAS = 2000
MAX_EDGES = 30000
REF = re.compile(r"(?:(?:'((?:[^']|'')+)'|([^'!]+))!)?(\$?[A-Z]{1,3}\$?[1-9][0-9]{0,6}(?::\$?[A-Z]{1,3}\$?[1-9][0-9]{0,6})?)", re.I)
STATIC_FUNCTIONS = set('SUM AVERAGE MIN MAX COUNT COUNTA COUNTIF COUNTIFS SUMIF SUMIFS AVERAGEIF AVERAGEIFS IF IFERROR IFNA AND OR NOT ABS ROUND ROUNDUP ROUNDDOWN INT MOD SQRT POWER PRODUCT SUMPRODUCT INDEX MATCH VLOOKUP HLOOKUP XLOOKUP XMATCH CONCAT CONCATENATE TEXT LEFT RIGHT MID LEN TRIM UPPER LOWER DATE YEAR MONTH DAY EOMONTH TODAY NOW ISERROR ISNUMBER ISTEXT'.split())


def label(sheet, cell):
    return "'" + sheet.replace("'", "''") + "'!" + cell


def references(formula, sheet, sheets):
    deps, issues = set(), []
    try:
        tokens = Tokenizer(formula).items
    except (TokenizerError, IndexError) as exc:
        return [], [{'kind': 'unresolved', 'reason': 'Formula tokenization failed'}]
    for token in tokens:
        value = token.value
        if token.type == 'FUNC' and token.subtype == 'OPEN':
            name = value[:-1].upper()
            if name not in STATIC_FUNCTIONS:
                issues.append({'kind': 'unresolved', 'reason': 'Dynamic or unsupported function', 'token': value})
        if token.type != 'OPERAND':
            continue
        if token.subtype == 'ERROR':
            issues.append({'kind': 'formula-error-token', 'reason': 'Error token in formula text', 'token': value})
        if token.subtype != 'RANGE':
            continue
        match = REF.fullmatch(value)
        if not match:
            issues.append({'kind': 'unresolved', 'reason': 'Unsupported reference (name, table, external, 3D or whole row/column)', 'token': value})
            continue
        quoted, plain, cells = match.groups()
        if any(c in (quoted if quoted is not None else plain or '') for c in '[]:'):
            issues.append({'kind': 'unresolved', 'reason': 'External or 3D reference unsupported', 'token': value})
            continue
        target = (quoted.replace("''", "'") if quoted is not None else plain) or sheet
        target = sheets.get(target.casefold())
        if target is None:
            issues.append({'kind': 'unresolved', 'reason': 'Referenced sheet is missing', 'token': value})
            continue
        c1, r1, c2, r2 = range_boundaries(cells.replace('$', '').upper())
        if c2 > 16384 or r2 > 1048576 or c1 > c2 or r1 > r2:
            issues.append({'kind': 'unresolved', 'reason': 'Invalid A1 bounds', 'token': value})
            continue
        if (c2-c1+1)*(r2-r1+1) > MAX_RANGE:
            issues.append({'kind': 'unresolved', 'reason': f'Range exceeds {MAX_RANGE} cells', 'token': value})
            continue
        deps.update(label(target, f'{get_column_letter(c)}{r}')
                    for r in range(r1, r2+1) for c in range(c1, c2+1))
        if len(deps) > MAX_EDGES:
            raise ValueError('Formula dependency limit exceeded')
    return sorted(deps), issues


def downstream(graph, seed):
    seen = {seed}; queue = deque([seed])
    while queue:
        for node in graph.get(queue.popleft(), []):
            if node not in seen:
                seen.add(node); queue.append(node)
    return sorted(seen - {seed})


def cycles(graph, nodes):
    """Iterative Kosaraju: cycles remain bounded even beyond Python recursion depth."""
    visited, order = set(), []
    for start in sorted(nodes):
        if start in visited: continue
        visited.add(start); stack = [(start, iter(graph.get(start, [])))]
        while stack:
            node, children = stack[-1]
            child = next(children, None)
            if child is None:
                order.append(node); stack.pop()
            elif child not in visited:
                visited.add(child); stack.append((child, iter(graph.get(child, []))))
    reverse = defaultdict(list)
    for a, bs in graph.items():
        for b in bs: reverse[b].append(a)
    visited = set(); groups = []
    for start in reversed(order):
        if start in visited: continue
        visited.add(start); stack = [start]; group = []
        while stack:
            node = stack.pop(); group.append(node)
            for child in reverse[node]:
                if child not in visited: visited.add(child); stack.append(child)
        if len(group) > 1 or start in graph.get(start, []): groups.append(sorted(group))
    return sorted(groups)


def inspect_workbook(data):
    if len(data) > 10_000_000: raise ValueError('Workbook limit: 10 MB compressed')
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        parts = archive.infolist()
        if len(parts) > 1000 or sum(p.file_size for p in parts) > 20_000_000:
            raise ValueError('Workbook archive exceeds inspection limits')
        if any(p.file_size > 8_000_000 for p in parts): raise ValueError('Workbook part exceeds 8 MB')
        if any('vbaproject' in p.filename.casefold() for p in parts): raise ValueError('Macro workbooks are unsupported')
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    try:
        if len(workbook.worksheets) > 30: raise ValueError('Workbook limit: 30 sheets')
        sheets = {s.title.casefold(): s.title for s in workbook}
        formulas, issues, sheet_info = {}, [], []
        total_cells, total_edges = 0, 0
        for ws in workbook:
            rows, cols = ws.max_row or 0, ws.max_column or 0
            if rows*cols > 200000: raise ValueError('Sheet used rectangle exceeds 200,000 cells')
            total_cells += rows*cols
            if total_cells > 500000: raise ValueError('Workbook used rectangles exceed 500,000 cells')
            sheet_info.append({'name': ws.title, 'visibility': ws.sheet_state})
            cached_rows = cached[ws.title].iter_rows()
            for row, cached_row in zip(ws.iter_rows(), cached_rows):
                for cell, cache in zip(row, cached_row):
                    if cell.data_type == 'e':
                        source = label(ws.title, cell.coordinate)
                        issues.append({'cell': source, 'kind': 'stored-error', 'reason': 'Stored error value', 'token': str(cell.value)})
                    if cell.data_type != 'f': continue
                    if len(formulas) >= MAX_FORMULAS: raise ValueError(f'Workbook limit: {MAX_FORMULAS} formulas')
                    source = label(ws.title, cell.coordinate)
                    if not isinstance(cell.value, str):
                        deps, found = [], [{'kind': 'unresolved', 'reason': 'Array or data-table formula unsupported'}]
                    else:
                        deps, found = references(cell.value, ws.title, sheets)
                    total_edges += len(deps)
                    if total_edges > MAX_EDGES: raise ValueError('Workbook dependency limit exceeded')
                    if cache.data_type == 'e':
                        found.append({'kind': 'cached-error', 'reason': 'Cached error; may be stale', 'token': str(cache.value)})
                    formulas[source] = {'sheet': ws.title, 'cell': cell.coordinate,
                                        'formula': cell.value if isinstance(cell.value, str) else None,
                                        'dependencies': deps, 'cache_present': cache.value is not None,
                                        'cached_error': str(cache.value) if cache.data_type == 'e' else None}
                    issues.extend({'cell': source, **issue} for issue in found)
        if len(issues) > 4000: raise ValueError('Workbook limit: 4,000 issue items')
        graph = defaultdict(set)
        for dependent, record in formulas.items():
            for precedent in record['dependencies']: graph[precedent].add(dependent)
        if sum(len(v) for v in graph.values()) > MAX_EDGES: raise ValueError('Workbook dependency limit exceeded')
        graph = {k: sorted(v) for k, v in sorted(graph.items())}
        cycle_groups = cycles(graph, set(formulas) | set(graph))
        seeds = sorted({i['cell'] for i in issues} | {n for g in cycle_groups for n in g})
        if len(seeds) > 2000: raise ValueError('Workbook limit: 2,000 issue seeds')
        impacts, reach_total = [], 0
        for seed in seeds:
            reachable = downstream(graph, seed)
            reach_total += len(reachable)
            if reach_total > 200000: raise ValueError('Workbook impact report exceeds 200,000 reachable entries')
            impacts.append({'cell': seed, 'reachable_formulas': reachable})
        impacts.sort(key=lambda r: (-len(r['reachable_formulas']), r['cell']))
        return {'schema_version': 1, 'input_sha256': hashlib.sha256(data).hexdigest(),
                'mode': 'static-dependency-audit', 'sheets': sheet_info, 'formulas': formulas,
                'issues': issues, 'cycles': cycle_groups, 'impacts': impacts, 'dependents': graph,
                'coverage': {'formula_cells': len(formulas), 'dependency_edges': sum(len(v) for v in graph.values()),
                             'unresolved_items': sum(i['kind']=='unresolved' for i in issues),
                             'formula_cells_without_cache': sum(not f['cache_present'] for f in formulas.values())},
                'interpretation': 'Edges are syntactic references, not proof of executed branches or wrong results. Unsupported references make impact counts incomplete. No formulas recalculated; cached errors may be stale.'}
    finally:
        workbook.close(); cached.close()


def build(path, output):
    from report import render
    if Path(path).suffix.lower() != '.xlsx': raise ValueError('Use a macro-free .xlsx workbook')
    plan = inspect_workbook(Path(path).read_bytes())
    output = Path(output); output.mkdir(mode=0o700)
    try:
        (output/'audit.json').write_text(json.dumps(plan, indent=2)+'\n', encoding='utf-8')
        render(plan, output)
        hashes = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in output.iterdir()}
        (output/'READY.json').write_text(json.dumps({'files': hashes}, indent=2)+'\n')
    except BaseException:
        shutil.rmtree(output); raise
    return plan


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path); parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args()
    try: result = build(args.workbook, args.out)
    except (ValueError, OSError, zipfile.BadZipFile) as exc: parser.exit(2, f'Audit failed: {exc}\n')
    print(json.dumps(result['coverage']))
