"""Generate a deliberately broken workbook and inspect its formula dependencies."""
import argparse
import hashlib
import json
from pathlib import Path
from openpyxl import Workbook
from audit import build


def make_fixture(path):
    # Deliberate defects exercise the audit; this is not a financial model.
    wb=Workbook(); wb.remove(wb.active)
    wb.properties.creator='SR Workbench'; wb.properties.lastModifiedBy='SR Workbench'
    for name in ['Summary','Sales','Costs','Rates','Checks','Notes']:
        wb.create_sheet(name)
    wb['Rates'].sheet_state='hidden'; wb['Checks'].sheet_state='veryHidden'
    wb['Rates']['A1']='Broken source'; wb['Rates']['B2']='=#REF!*1.05'
    for name in ['Sales','Costs']:
        ws=wb[name]; ws.append(['Illustrative period','Formula'])
        for n in range(2,6):
            ws.cell(n,1,n-1); ws.cell(n,2,f"='Rates'!$B$2*{n}")
    wb['Summary']['A1']='Dependency audit example'
    wb['Summary']['B2']="=SUM('Sales'!B2:B5)"
    wb['Summary']['B3']="=SUM('Costs'!B2:B5)"
    wb['Summary']['B4']='=B2-B3'
    wb['Checks']['A1']="='Summary'!B4"
    wb['Notes']['A1']='=B1'; wb['Notes']['B1']='=A1'
    wb['Notes']['A3']="='[old-book.xlsx]Archive'!A1"
    wb['Notes']['A4']='=INDIRECT("Summary!B4")'
    wb['Notes']['A5']='=UnresolvedRate*2'
    for ws in wb:
        ws.column_dimensions['A'].width=32; ws.column_dimensions['B'].width=30
    wb.save(path)


def run(output):
    output.mkdir(mode=0o700)
    source=output/'model.xlsx'; make_fixture(source)
    digest=hashlib.sha256(source.read_bytes()).hexdigest()
    plan=build(source,output/'report')
    assert plan['impacts'][0]['cell']=="'Rates'!B2"
    assert len(plan['impacts'][0]['reachable_formulas'])==12
    assert plan['cycles']==[["'Notes'!A1", "'Notes'!B1"]]
    assert plan['coverage']['unresolved_items']==3
    assert digest==hashlib.sha256(source.read_bytes()).hexdigest()
    proof={'verified_demo':True,'source_unchanged':True,'potentially_affected_formulas':12,
           'cycle_groups':1,'unresolved_items':3,'coverage':plan['coverage']}
    (output/'demo-proof.json').write_text(json.dumps(proof,indent=2)+'\n')
    return proof


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--out',type=Path,required=True)
    print(json.dumps(run(p.parse_args().out)))
