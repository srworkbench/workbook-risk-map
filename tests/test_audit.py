import hashlib
import io
import json
from pathlib import Path
import tempfile
import unittest
import zipfile
from unittest.mock import patch
from openpyxl import Workbook
from audit import references, downstream, cycles, inspect_workbook, build
from demo import run, make_fixture


class AuditTests(unittest.TestCase):
    def test_quoted_apostrophe_case_and_absolute_range(self):
        deps, issues=references("=SUM('O''Brien'!$B2:C$3)+a1",'Main',{'main':'Main',"o'brien":"O'Brien"})
        self.assertFalse(issues)
        self.assertEqual(deps,["'Main'!A1","'O''Brien'!B2","'O''Brien'!B3","'O''Brien'!C2","'O''Brien'!C3"])

    def test_strings_do_not_create_edges_or_error_tokens(self):
        deps, issues=references('=IF(A1="B2", "#REF!", "C3")','Main',{'main':'Main'})
        self.assertEqual(deps,["'Main'!A1"]);self.assertEqual(issues,[])

    def test_unsupported_references_are_visible(self):
        for formula in ['=NamedRange','=Table1[Amount]',"='[book.xlsx]Main'!A1",'=SUM(A:A)',
                        '=SUM(A1:A3000)','=INDIRECT("A1")','=OFFSET(A1,1,1)',"='Missing'!A1",'=SUM(First:Last!A1)']:
            with self.subTest(formula=formula):
                deps,issues=references(formula,'Main',{'main':'Main'})
                self.assertTrue(any(i['kind']=='unresolved' for i in issues))

    def test_error_literal_is_distinct_from_unresolved(self):
        deps,issues=references('=#REF!+A1','Main',{'main':'Main'})
        self.assertEqual(deps,["'Main'!A1"])
        self.assertEqual(issues[0]['kind'],'formula-error-token')

    def test_branch_union_is_potential_reachability(self):
        graph={'A':['B','C'],'B':['D'],'C':['D'],'D':['E'],'E':['D']}
        self.assertEqual(downstream(graph,'A'),['B','C','D','E'])
        self.assertEqual(cycles(graph,set(graph)),[['D','E']])
        self.assertEqual(downstream(graph,'D'),['E'])

    def test_long_cycle_and_self_cycle_do_not_recurse(self):
        graph={str(n):[str((n+1)%1500)] for n in range(1500)}
        self.assertEqual(len(cycles(graph,set(graph))[0]),1500)
        self.assertEqual(cycles({'A':['A']},{'A'}),[['A']])

    def test_bad_archive_and_size_limit(self):
        with self.assertRaises(zipfile.BadZipFile): inspect_workbook(b'not a zip')
        with self.assertRaisesRegex(ValueError,'10 MB'): inspect_workbook(b'0'*10000001)

    def test_dimension_limit_refuses_before_large_iteration(self):
        wb=Workbook();wb.active['XFD1000']='x';data=io.BytesIO();wb.save(data)
        with self.assertRaisesRegex(ValueError,'rectangle'): inspect_workbook(data.getvalue())

    def test_stored_error_reaches_dependent_and_cache_is_not_evaluated(self):
        wb=Workbook();ws=wb.active;ws.title='Main';ws['A1']='#NAME?';ws['B1']='=A1';data=io.BytesIO();wb.save(data)
        plan=inspect_workbook(data.getvalue())
        self.assertEqual(plan['issues'][0]['kind'],'stored-error')
        self.assertEqual(plan['impacts'][0]['reachable_formulas'],["'Main'!B1"])
        self.assertFalse(plan['formulas']["'Main'!B1"]['cache_present'])

    def test_real_demo_and_immutable_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            out=Path(tmp)/'demo'; proof=run(out); self.assertTrue(proof['verified_demo'])
            report=out/'report'; plan=json.loads((report/'audit.json').read_text())
            states={s['name']:s['visibility'] for s in plan['sheets']}
            self.assertEqual(states['Rates'],'hidden');self.assertEqual(states['Checks'],'veryHidden')
            original=(report/'audit.json').read_bytes()
            with self.assertRaises(FileExistsError): build(out/'model.xlsx',report)
            self.assertEqual(original,(report/'audit.json').read_bytes())
            for name,digest in json.loads((report/'READY.json').read_text())['files'].items():
                self.assertEqual(hashlib.sha256((report/name).read_bytes()).hexdigest(),digest)

    def test_cached_error_is_reported_as_stale_not_recalculated(self):
        wb=Workbook();wb.active['A1']='=1/0';data=io.BytesIO();wb.save(data)
        revised=io.BytesIO()
        with zipfile.ZipFile(data) as src, zipfile.ZipFile(revised,'w') as dst:
            for name in src.namelist():
                blob=src.read(name)
                if name=='xl/worksheets/sheet1.xml':
                    blob=blob.replace(b'<c r="A1">',b'<c r="A1" t="e">').replace(b'<v></v>',b'<v>#DIV/0!</v>')
                dst.writestr(name,blob)
        plan=inspect_workbook(revised.getvalue())
        self.assertEqual(plan['issues'][0]['kind'],'cached-error')
        self.assertIn('stale',plan['issues'][0]['reason'])

    def test_no_issues_is_not_a_clean_bill_of_health(self):
        wb=Workbook();wb.active['A1']='=1/0';data=io.BytesIO();wb.save(data)
        plan=inspect_workbook(data.getvalue())
        self.assertEqual(plan['impacts'],[])
        self.assertEqual(plan['coverage']['formula_cells_without_cache'],1)
        self.assertIn('No formulas recalculated',plan['interpretation'])

    def test_issue_seed_budget(self):
        wb=Workbook()
        for row in range(1,2002): wb.active.cell(row,1,'#REF!')
        data=io.BytesIO();wb.save(data)
        with self.assertRaisesRegex(ValueError,'issue seeds'): inspect_workbook(data.getvalue())

    def test_render_failure_cleanup_and_retry(self):
        with tempfile.TemporaryDirectory() as tmp:
            source=Path(tmp)/'model.xlsx';make_fixture(source);out=Path(tmp)/'out'
            with patch('report.render',side_effect=RuntimeError('injected failure')):
                with self.assertRaises(RuntimeError): build(source,out)
            self.assertFalse(out.exists());build(source,out);self.assertTrue((out/'READY.json').exists())


if __name__=='__main__': unittest.main()
